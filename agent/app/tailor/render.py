"""Uyarlanmış CV'yi PDF ve DOCX'e basar.

Font notu: fpdf2'nin gömülü fontları Latin-1 ile sınırlı, yani "ğ ş İ ı"
karakterlerini basamıyor — Türkçe bir CV'de bu kabul edilemez. Bu yüzden
sistemdeki DejaVuSans TTF'ini gömüyoruz. Font bulunamazsa PDF üretimini
sessizce bozuk harflerle sürdürmek yerine açık hata veriyoruz.
"""

from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from fpdf import FPDF

from ..schemas import TailoredCV

#: Türkçe karakterleri kapsayan, çoğu dağıtımda hazır gelen font
_FONT_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    "/usr/share/fonts/TTF/DejaVuSans.ttf",
    "/Library/Fonts/DejaVuSans.ttf",
)
_FONT_BOLD_CANDIDATES = (
    "/usr/share/fonts/truetype/dejavu/DejaVuSans-Bold.ttf",
    "/usr/share/fonts/TTF/DejaVuSans-Bold.ttf",
    "/Library/Fonts/DejaVuSans-Bold.ttf",
)

INK = (17, 24, 39)
MUTED = (107, 114, 128)
RULE = (209, 213, 219)


#: Bölüm başlıkları. Kit ilanın dilinde üretiliyor; İngilizce bir CV'de
#: "ÖZET / DENEYİM" yazması belgeyi bozuk gösterir.
_HEADINGS = {
    "tr": {
        "summary": "Özet", "skills": "Yetkinlikler", "experience": "Deneyim",
        "education": "Eğitim", "languages": "Diller", "page": "sayfa",
    },
    "en": {
        "summary": "Summary", "skills": "Skills", "experience": "Experience",
        "education": "Education", "languages": "Languages", "page": "page",
    },
}


def tr_upper(text: str) -> str:
    """Türkçe kurallarına göre büyük harfe çevirir.

    Python'un str.upper() metodu "i" -> "I" yapıyor; Türkçe'de doğrusu
    "i" -> "İ". Bölüm başlıklarında bu "YETKINLIKLER" / "EĞITIM" gibi
    hatalara yol açıyor ve işverene giden belgede göze batıyor.
    """
    return text.replace("i", "İ").replace("ı", "I").upper()


def upper_for(language: str, text: str) -> str:
    """Dile uygun büyük harf. İngilizce metne Türkçe kuralı uygulamak
    "Skills" -> "SKİLLS" gibi bozuk çıktı üretir."""
    return tr_upper(text) if language == "tr" else text.upper()


def labels(language: str) -> dict[str, str]:
    return _HEADINGS.get(language, _HEADINGS["tr"])


class FontMissing(RuntimeError):
    pass


def _find_font(candidates: tuple[str, ...]) -> Path:
    for path in candidates:
        p = Path(path)
        if p.exists():
            return p
    raise FontMissing(
        "Türkçe karakterleri basabilen DejaVuSans fontu bulunamadı. "
        "Kurulum: sudo apt install fonts-dejavu-core"
    )


class _CVPdf(FPDF):
    def __init__(self, name: str, language: str = "tr") -> None:
        super().__init__(format="A4", unit="mm")
        self._candidate = name
        self._language = language
        self.set_auto_page_break(auto=True, margin=16)
        self.set_margins(18, 16, 18)
        self.add_font("DejaVu", "", str(_find_font(_FONT_CANDIDATES)))
        self.add_font("DejaVu", "B", str(_find_font(_FONT_BOLD_CANDIDATES)))

    def footer(self) -> None:
        self.set_y(-12)
        self.set_font("DejaVu", "", 7.5)
        self.set_text_color(*MUTED)
        page_word = labels(self._language)["page"]
        self.cell(0, 5, f"{self._candidate} · {page_word} {self.page_no()}", align="C")

    # -- yardımcılar ------------------------------------------------------
    def heading(self, text: str) -> None:
        self.ln(3)
        self.set_font("DejaVu", "B", 10.5)
        self.set_text_color(*INK)
        self.cell(0, 6, upper_for(self._language, text), new_x="LMARGIN", new_y="NEXT")
        self.set_draw_color(*RULE)
        self.set_line_width(0.3)
        y = self.get_y()
        self.line(self.l_margin, y, self.w - self.r_margin, y)
        self.ln(2)

    def body(self, text: str, size: float = 9.5, color: tuple = INK) -> None:
        self.set_font("DejaVu", "", size)
        self.set_text_color(*color)
        self.multi_cell(0, 4.8, text, new_x="LMARGIN", new_y="NEXT")

    def bullet(self, text: str) -> None:
        self.set_font("DejaVu", "", 9.5)
        self.set_text_color(*INK)
        left = self.get_x()
        self.cell(4, 4.8, "•")
        self.set_x(left + 4)
        self.multi_cell(0, 4.8, text, new_x="LMARGIN", new_y="NEXT")


def to_pdf(cv: TailoredCV, language: str = "tr") -> bytes:
    lab = labels(language)
    pdf = _CVPdf(cv.full_name, language)
    pdf.add_page()

    pdf.set_font("DejaVu", "B", 17)
    pdf.set_text_color(*INK)
    pdf.cell(0, 9, cv.full_name, new_x="LMARGIN", new_y="NEXT")

    if cv.headline:
        pdf.set_font("DejaVu", "", 10.5)
        pdf.set_text_color(*MUTED)
        pdf.multi_cell(0, 5, cv.headline, new_x="LMARGIN", new_y="NEXT")
    if cv.contact:
        pdf.set_font("DejaVu", "", 8.5)
        pdf.set_text_color(*MUTED)
        pdf.multi_cell(0, 4.5, cv.contact, new_x="LMARGIN", new_y="NEXT")

    if cv.summary:
        pdf.heading(lab["summary"])
        pdf.body(cv.summary)

    if cv.skills:
        pdf.heading(lab["skills"])
        pdf.body(" · ".join(cv.skills))

    if cv.experience:
        pdf.heading(lab["experience"])
        for exp in cv.experience:
            pdf.set_font("DejaVu", "B", 10)
            pdf.set_text_color(*INK)
            pdf.multi_cell(0, 5, f"{exp.title} — {exp.company}", new_x="LMARGIN", new_y="NEXT")
            if exp.period:
                pdf.set_font("DejaVu", "", 8.5)
                pdf.set_text_color(*MUTED)
                pdf.cell(0, 4.5, exp.period, new_x="LMARGIN", new_y="NEXT")
            for line in exp.bullets:
                pdf.bullet(line)
            pdf.ln(1.5)

    if cv.education:
        pdf.heading(lab["education"])
        for line in cv.education:
            pdf.bullet(line)

    if cv.languages:
        pdf.heading(lab["languages"])
        pdf.body(" · ".join(cv.languages))

    return bytes(pdf.output())


def to_docx(cv: TailoredCV, language: str = "tr") -> bytes:
    """DOCX çıktısı — bazı başvuru sistemleri PDF kabul etmiyor."""
    lab = labels(language)
    doc = Document()
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style.font.size = Pt(10)

    name = doc.add_paragraph()
    run = name.add_run(cv.full_name)
    run.bold = True
    run.font.size = Pt(20)

    if cv.headline:
        p = doc.add_paragraph()
        r = p.add_run(cv.headline)
        r.font.size = Pt(11)
        r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
    if cv.contact:
        p = doc.add_paragraph()
        r = p.add_run(cv.contact)
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)

    def section(title: str) -> None:
        h = doc.add_paragraph()
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(2)
        r = h.add_run(upper_for(language, title))
        r.bold = True
        r.font.size = Pt(10)

    if cv.summary:
        section(lab["summary"])
        doc.add_paragraph(cv.summary)

    if cv.skills:
        section(lab["skills"])
        doc.add_paragraph(" · ".join(cv.skills))

    if cv.experience:
        section(lab["experience"])
        for exp in cv.experience:
            p = doc.add_paragraph()
            r = p.add_run(f"{exp.title} — {exp.company}")
            r.bold = True
            if exp.period:
                sub = doc.add_paragraph()
                sr = sub.add_run(exp.period)
                sr.font.size = Pt(9)
                sr.font.color.rgb = RGBColor(0x6B, 0x72, 0x80)
            for line in exp.bullets:
                doc.add_paragraph(line, style="List Bullet")

    if cv.education:
        section(lab["education"])
        for line in cv.education:
            doc.add_paragraph(line, style="List Bullet")

    if cv.languages:
        section(lab["languages"])
        doc.add_paragraph(" · ".join(cv.languages))

    for p in doc.paragraphs:
        p.alignment = WD_ALIGN_PARAGRAPH.LEFT

    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


#: Content-Disposition başlığı latin-1 ile kodlanıyor; "ş" gibi harfler orada
#: yok ve indirme 500 ile patlıyor. Bu yüzden dosya adını ASCII'ye çeviriyoruz.
_ASCII_MAP = str.maketrans(
    {
        "ı": "i", "İ": "I", "ş": "s", "Ş": "S", "ğ": "g", "Ğ": "G",
        "ü": "u", "Ü": "U", "ö": "o", "Ö": "O", "ç": "c", "Ç": "C",
        "â": "a", "î": "i", "û": "u", "é": "e", "ß": "ss",
    }
)


def safe_filename(candidate: str, company: str, extension: str) -> str:
    """İndirilen dosyaya okunabilir, ASCII ve güvenli bir ad ver."""
    parts = [p for p in (candidate, company) if p]
    raw = "_".join(parts) or "CV"
    ascii_only = raw.translate(_ASCII_MAP).encode("ascii", "ignore").decode("ascii")
    cleaned = "".join(c if c.isalnum() or c in " -_" else "" for c in ascii_only)
    stem = "_".join(cleaned.split())[:80].strip("_")
    return f"{stem}_CV.{extension}" if stem else f"CV.{extension}"
